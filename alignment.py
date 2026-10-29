import os
import re
import openpyxl
from config import EMAIL_DOMAIN, EMAIL_OVERRIDES


def region_key(value):
    """Normalize any region label to its parent region number.
    '01' -> '1', '30.1' -> '30', 'Sales Region 25' -> '25', 25.2 -> '25'."""
    s = str(value).strip()
    m = re.match(r"(?:Sales Region\s+)?0*(\d+)", s, re.IGNORECASE)
    if not m:
        return s
    return m.group(1) or "0"


def _clean(val):
    if val is None:
        return None
    s = str(val).strip()
    if s in ("", "#"):
        return None
    return s


def split_names(cell):
    """'Peter Ungaro/Jessica Fryger' -> ['Peter Ungaro', 'Jessica Fryger']."""
    s = _clean(cell)
    if not s:
        return []
    return [p.strip() for p in re.split(r"[/;,]", s) if p.strip()]


def first_name(full_name):
    parts = str(full_name).strip().split()
    return parts[0] if parts else ""


def make_email(full_name):
    """First.Last@loves.com from a full name. Returns None if a first and last name can't be found."""
    key = " ".join(str(full_name or "").split()).lower()
    for name, addr in EMAIL_OVERRIDES.items():
        if " ".join(name.split()).lower() == key:
            return addr
    cleaned = re.sub(r"[^A-Za-z\-\s]", "", str(full_name or ""))
    parts = cleaned.split()
    if len(parts) < 2:
        return None
    return f"{parts[0]}.{parts[-1]}@{EMAIL_DOMAIN}"


def is_open_seat(name):
    return (name or "").strip().upper().startswith("OPEN")


def parse_alignment(path):
    """Read the FUEL tab of the Region Alignments workbook.

    Returns {
      "regions": {code: {"code","label","rep","area_managers":[...],"group","manager","director"}},
      "groups":  {"National": {codes}, "East": {...}, "West": {...}, "Inside East": {...}, ...}
    }
    """
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet = next((s for s in wb.sheetnames if s.strip().upper() == "FUEL"), None)
    if sheet is None:
        raise ValueError("Alignment file has no 'FUEL' tab.")
    rows = list(wb[sheet].iter_rows(values_only=True))
    if not rows:
        raise ValueError("Alignment file 'FUEL' tab is empty.")

    header = [str(c).strip() if c is not None else "" for c in rows[0]]

    def col(name):
        if name not in header:
            raise ValueError(f"Alignment file 'FUEL' tab is missing the '{name}' column.")
        return header.index(name)

    c_region, c_rep = col("Region"), col("Sales Person")
    c_am = col("Area Manager")
    c_mgr = header.index("Manager") if "Manager" in header else None
    c_dir = header.index("Director") if "Director" in header else None

    regions, groups = {}, {}
    current_group = None
    for row in rows[1:]:
        label = _clean(row[c_region]) if len(row) > c_region else None
        rep = _clean(row[c_rep]) if len(row) > c_rep else None
        if not label:
            continue
        if rep == "Reporting Group":
            current_group = label
            groups.setdefault(label, set())
            continue
        m = re.match(r"Sales Region\s+(\d+)", label, re.IGNORECASE)
        if not m:
            continue
        code = m.group(1).lstrip("0") or "0"
        regions[code] = {
            "code": code,
            "label": label,
            "rep": rep,
            "area_managers": split_names(row[c_am]) if len(row) > c_am else [],
            "manager": _clean(row[c_mgr]) if c_mgr is not None and len(row) > c_mgr else None,
            "director": _clean(row[c_dir]) if c_dir is not None and len(row) > c_dir else None,
            "group": current_group,
        }
        if current_group:
            groups[current_group].add(code)
    if not regions:
        raise ValueError("No 'Sales Region ##' rows were found on the alignment file's 'FUEL' tab.")
    return {"regions": regions, "groups": groups}


def inside_region_codes(alignment):
    """Union of every group whose name starts with 'Inside' (Inside East + Inside West)."""
    codes = set()
    for name, members in alignment["groups"].items():
        if name.strip().lower().startswith("inside"):
            codes |= members
    return codes


SAVED_ALIGNMENT_NAME = "Region_Alignments.xlsx"


def find_saved_alignment(base_dir):
    """Path to the alignment workbook saved next to the app, or None.
    Prefers Region_Alignments.xlsx; otherwise accepts exactly one *alignment*.xlsx in the folder.
    Raises ValueError if several candidates exist (so the wrong one is never used silently)."""
    exact = os.path.join(base_dir, SAVED_ALIGNMENT_NAME)
    if os.path.isfile(exact):
        return exact
    matches = sorted(
        f for f in os.listdir(base_dir)
        if "alignment" in f.lower() and f.lower().endswith(".xlsx") and not f.startswith("~$")
    )
    if len(matches) == 1:
        return os.path.join(base_dir, matches[0])
    if len(matches) > 1:
        raise ValueError(
            "More than one alignment file is saved with the app (" + ", ".join(matches)
            + "). Keep only one, named " + SAVED_ALIGNMENT_NAME + "."
        )
    return None
