import openpyxl
from config import (
    TTC_METRIC_SHEETS, TTC_CUSTOMER_COLS, TTC_REGION_COLS, TTC_WEEK_COUNT,
    TTC_KNOWN_MANAGERS, TTC_INSIDE_MANAGER, TTC_CUSTOMER_DATA_START_ROW,
)


def safe_float(val):
    if val is None or val in ("", "%", "#"):
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def clean_field(val):
    if val is None or val == "#" or str(val).strip() == "#":
        return None
    s = str(val).strip()
    return s if s else None


def get_week_dates(ws):
    """Pull the 13 week-ending dates from the header row, newest first."""
    rows = list(ws.iter_rows(values_only=True))
    header = rows[4]
    dates = []
    for cell in header:
        if cell is None:
            continue
        text = str(cell)
        last_part = text.split("\n")[-1].strip()
        if "/" in last_part and len(last_part) <= 5:
            dates.append(last_part)
    return dates[-TTC_WEEK_COUNT:] if len(dates) >= TTC_WEEK_COUNT else dates


def parse_ttc_roster(ws):
    """Build rep -> senior manager mapping from the Region file's hierarchy block.
    Row order is NOT reliable (varies by sheet), so managers are detected by name,
    and 'Inside' is excluded entirely regardless of where it falls in the sheet."""
    rows = list(ws.iter_rows(values_only=True))
    start = None
    for i, row in enumerate(rows):
        if row and row[0] == "Truck Care Hierarchy":
            start = i
            break
    if start is None:
        return {}

    roster = {}
    current_manager = None
    for i in range(start + 1, len(rows)):
        name = rows[i][0] if rows[i] else None
        if name is None or str(name).strip() == "":
            break
        name = str(name).strip()
        if name in TTC_KNOWN_MANAGERS:
            current_manager = name
            continue
        if name == "Inside":
            current_manager = None  # Inside Sales has its own separate manager -- exclude entirely, don't attribute to Brett/Keith/TJ
            continue
        if current_manager:
            roster[name] = current_manager
    return roster


def build_master_roster(wb_region):
    """Merge the roster across all 4 metric sheets for maximum coverage."""
    roster = {}
    for sheet in TTC_METRIC_SHEETS:
        if sheet in wb_region.sheetnames:
            for rep, mgr in parse_ttc_roster(wb_region[sheet]).items():
                roster[rep] = mgr
    return roster


def get_manager_rollup_row(ws, manager):
    """Pull the senior manager's own subtotal row (already aggregated by the report)."""
    rows = list(ws.iter_rows(values_only=True))
    for row in rows:
        if row and row[0] == manager:
            return row
    return None


def extract_week_series(row, col_map):
    start = col_map["week_start"]
    values = []
    for i in range(TTC_WEEK_COUNT):
        idx = start + i
        v = safe_float(row[idx]) if idx < len(row) else None
        values.append(v)
    return values


def analyze_manager_rollup(ws, sheet_name, manager, week_dates):
    """This week / 13wk avg / curr-vs-avg / YOY / 13-week high-low, straight from the
    manager's own subtotal row in the Region file -- no re-aggregation needed."""
    col_map = TTC_REGION_COLS[sheet_name]
    row = get_manager_rollup_row(ws, manager)
    if row is None:
        return None

    avg_13wk = safe_float(row[col_map["avg_13wk"]])
    curr_ov_avg = safe_float(row[col_map["curr_ov_avg"]])
    yoy = safe_float(row[col_map["yoy"]])
    week_values = extract_week_series(row, col_map)

    this_week = week_values[0] if week_values else None
    valid_pairs = [(d, v) for d, v in zip(week_dates, week_values) if v is not None]
    if not valid_pairs:
        return None
    high_week, high_val = max(valid_pairs, key=lambda x: x[1])
    low_week, low_val = min(valid_pairs, key=lambda x: x[1])

    return {
        "manager": manager,
        "this_week": this_week,
        "avg_13wk": avg_13wk,
        "curr_ov_avg": curr_ov_avg,
        "yoy": yoy,
        "high_val": high_val, "high_week": high_week,
        "low_val": low_val, "low_week": low_week,
        "week_series": list(zip(week_dates, week_values)),
    }


def analyze_rep_swings(ws, sheet_name, manager, roster):
    """This-week-vs-prior-week delta for every rep under this manager -- used to find
    the single biggest unit-count swing, which is the real 'story' (the team rollup's
    own week-over-week move is often too small to be meaningful on its own).
    Also pulls last year's same-week value for that rep, computed from the sheet's
    own YOY-delta column (the source file stores the DIFFERENCE vs last year, not
    last year's absolute value, so last_year = this_week - yoy_delta)."""
    col_map = TTC_REGION_COLS[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    swings = []
    for row in rows:
        if not row or not row[0]:
            continue
        name = str(row[0]).strip()
        if name in TTC_KNOWN_MANAGERS or name == "Inside" or name == "Truck Care Hierarchy":
            continue
        if roster.get(name) != manager:
            continue
        week_start = col_map["week_start"]
        this_week = safe_float(row[week_start]) if week_start < len(row) else None
        prior_week = safe_float(row[week_start + 1]) if week_start + 1 < len(row) else None
        yoy_delta = safe_float(row[col_map["yoy"]]) if col_map["yoy"] < len(row) else None
        if this_week is None or prior_week is None:
            continue
        delta = this_week - prior_week
        last_year = (this_week - yoy_delta) if yoy_delta is not None else None
        swings.append({"name": name, "this_week": this_week, "prior_week": prior_week, "delta": delta, "last_year": last_year})
    swings.sort(key=lambda x: abs(x["delta"]), reverse=True)
    return swings


def analyze_ttc_customers(ws, manager, roster):
    """Per-customer rows for one manager, one metric sheet -- analogous to the fuel
    analyzer's per-account records, tagged by manager via the roster lookup."""
    col_map = TTC_CUSTOMER_COLS
    rows = list(ws.iter_rows(values_only=True))
    records = []
    for i, row in enumerate(rows):
        if i < TTC_CUSTOMER_DATA_START_ROW:
            continue
        rep = clean_field(row[col_map["rep"]])
        customer = clean_field(row[col_map["customer"]])
        if not customer or customer == "Customer":
            continue
        if not rep or roster.get(rep) != manager:
            continue
        salesperson = clean_field(row[col_map["salesperson"]])
        area_manager = clean_field(row[col_map["area_manager"]])
        avg_13wk = safe_float(row[col_map["avg_13wk"]])
        yoy = safe_float(row[col_map["yoy"]])
        week_values = extract_week_series(row, col_map)
        this_week = week_values[0] if week_values else None
        prior_week = week_values[1] if len(week_values) > 1 else None
        delta = (this_week - prior_week) if (this_week is not None and prior_week is not None) else None
        records.append({
            "customer": customer, "rep": rep, "salesperson": salesperson,
            "area_manager": area_manager, "avg_13wk": avg_13wk, "yoy": yoy,
            "this_week": this_week, "prior_week": prior_week, "delta": delta,
        })
    return records


def find_col(header_row, label, exclude=None):
    """Locate a Monthly-file column by its header text rather than a fixed index --
    the National tab has one fewer metadata column than East/West, so fixed indices
    silently misalign between tabs."""
    for i, cell in enumerate(header_row):
        if cell is None:
            continue
        text = str(cell).replace("\n", " ").strip()
        if label in text:
            if exclude and exclude in text:
                continue
            return i
    return None


def parse_monthly_history_tab(ws, roster, month_start_idx, known_managers):
    """A '13 Mo History' tab is customer-grain, wide-format: rep | customer | SAP ID |
    region code | (category/unit) | 13 month columns (newest first) | Overall Result.
    Sums every customer row up to its manager via the SAME roster built from the
    weekly Region file -- rep names are consistent between the weekly and monthly
    exports, so no separate monthly-specific roster is needed."""
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]
    month_labels = [header[month_start_idx + i] for i in range(13)]

    manager_totals = {m: [0.0] * 13 for m in known_managers}
    manager_has_data = {m: [False] * 13 for m in known_managers}

    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        rep = str(row[0]).strip()
        manager = roster.get(rep)
        if manager not in known_managers:
            continue
        for i in range(13):
            idx = month_start_idx + i
            val = row[idx] if idx < len(row) else None
            if val is not None:
                try:
                    manager_totals[manager][i] += float(val)
                    manager_has_data[manager][i] = True
                except (TypeError, ValueError):
                    pass

    return month_labels, manager_totals, manager_has_data


def compute_monthly_high_low(month_labels, series, has_data):
    valid = [(month_labels[i], series[i]) for i in range(13) if has_data[i]]
    if not valid:
        return None
    high_month, high_val = max(valid, key=lambda x: x[1])
    low_month, low_val = min(valid, key=lambda x: x[1])
    return {
        "this_month": series[0],
        "last_year": series[12] if has_data[12] else None,
        "high_val": high_val, "high_month": high_month,
        "low_val": low_val, "low_month": low_month,
    }


def analyze_ttc_monthly(monthly_path, roster=None):
    """Pull Profit/PPG and the new-vs-retread tire split for each manager's tab
    (current-month snapshot), plus a true 13-month trend (high/low, this month,
    same month last year) for Tires & Retreads, PM, and Labor Hrs, if the uploaded
    file has those '13 Mo History' tabs. TCE Spend per Truck has no equivalent
    history tab in this report, so it stays snapshot-only. Entirely optional data --
    only called if a Monthly file was uploaded. roster is the same manager-roster
    already built from the weekly Region file."""
    from config import TTC_MONTHLY_TAB_BY_MANAGER, TTC_KNOWN_MANAGERS

    results = {"period": None, "by_manager": {}, "monthly_trend": {}, "errors": []}
    roster = roster or {}
    try:
        wb = openpyxl.load_workbook(monthly_path, data_only=True)
        for i, row in enumerate(wb[list(wb.sheetnames)[1]].iter_rows(values_only=True)):
            if i == 2 and row[1]:
                results["period"] = str(row[1])
                break

        for manager, tab in TTC_MONTHLY_TAB_BY_MANAGER.items():
            if tab not in wb.sheetnames:
                continue
            ws = wb[tab]
            rows = list(ws.iter_rows(values_only=True))
            header = rows[6]
            grand_total = rows[8]

            col_profit = find_col(header, "Profit [All-In]", exclude="12 Month")
            col_ppg = find_col(header, "PPG [All-In]", exclude="12 Month")
            col_tireqty = find_col(header, "Tire")
            col_retreadqty = find_col(header, "Retread")
            col_yoy_tires = find_col(header, "YOY Tires")

            results["by_manager"][manager] = {
                "profit_all_in": safe_float(grand_total[col_profit]) if col_profit is not None else None,
                "ppg_all_in": safe_float(grand_total[col_ppg]) if col_ppg is not None else None,
                "tire_qty": safe_float(grand_total[col_tireqty]) if col_tireqty is not None else None,
                "retread_qty": safe_float(grand_total[col_retreadqty]) if col_retreadqty is not None else None,
                "yoy_tires_retreads": safe_float(grand_total[col_yoy_tires]) if col_yoy_tires is not None else None,
            }

        history_tabs = {
            "Tires & Retreads": ("13 Mo History -Tires & Retreads", 6),
            "PM": ("13 Mo History - PM", 5),
            "Labor Hrs": ("13 Mo History - LMS Labor Hrs", 5),
        }
        for metric, (sheet_name, month_start) in history_tabs.items():
            if sheet_name not in wb.sheetnames:
                continue
            ws = wb[sheet_name]
            month_labels, totals, has_data = parse_monthly_history_tab(
                ws, roster, month_start, TTC_KNOWN_MANAGERS
            )
            results["monthly_trend"][metric] = {"month_labels": month_labels, "by_manager": {}}
            for manager in TTC_KNOWN_MANAGERS:
                hl = compute_monthly_high_low(month_labels, totals[manager], has_data[manager])
                if hl:
                    hl["series"] = totals[manager]
                    hl["has_data"] = has_data[manager]
                    results["monthly_trend"][metric]["by_manager"][manager] = hl
    except Exception as e:
        results["errors"].append(f"Monthly file error: {str(e)}")
    return results


def filter_significant_customers(records, metric, thresholds):
    """Keep only customers whose this-week change clears the per-metric threshold --
    mirrors the fuel report's gallon-threshold filtering so the narrative isn't
    drowned in thousands of near-zero-change accounts."""
    threshold = thresholds.get(metric, 0)
    flagged = [r for r in records if r["delta"] is not None and abs(r["delta"]) >= threshold]
    flagged.sort(key=lambda x: abs(x["delta"]), reverse=True)
    return flagged


def analyze_ttc_reports(customer_path, region_path, monthly_path=None):
    """Top-level entry point, mirrors analyze_reports() for fuel."""
    from config import TTC_ALERT_THRESHOLDS

    results = {
        "week_dates": [],
        "roster": {},
        "manager_rollups": {},    # manager -> {sheet -> rollup dict}
        "manager_swings": {},     # manager -> {sheet -> sorted list of rep swings}
        "customer_records": {},   # manager -> {sheet -> list of ALL customer records}
        "flagged_customers": {},  # manager -> {sheet -> list of records clearing the threshold}
        "monthly": None,
        "errors": [],
    }

    try:
        wb_region = openpyxl.load_workbook(region_path, data_only=True)
        if "Tires" in wb_region.sheetnames:
            results["week_dates"] = get_week_dates(wb_region["Tires"])
        results["roster"] = build_master_roster(wb_region)

        for manager in TTC_KNOWN_MANAGERS:
            results["manager_rollups"][manager] = {}
            results["manager_swings"][manager] = {}
            for sheet in TTC_METRIC_SHEETS:
                if sheet not in wb_region.sheetnames:
                    continue
                ws = wb_region[sheet]
                rollup = analyze_manager_rollup(ws, sheet, manager, results["week_dates"])
                if rollup:
                    results["manager_rollups"][manager][sheet] = rollup
                swings = analyze_rep_swings(ws, sheet, manager, results["roster"])
                results["manager_swings"][manager][sheet] = swings
    except Exception as e:
        results["errors"].append(f"Region file error: {str(e)}")

    try:
        wb_customer = openpyxl.load_workbook(customer_path, data_only=True)
        for manager in TTC_KNOWN_MANAGERS:
            results["customer_records"][manager] = {}
            results["flagged_customers"][manager] = {}
            for sheet in TTC_METRIC_SHEETS:
                if sheet not in wb_customer.sheetnames:
                    continue
                ws = wb_customer[sheet]
                records = analyze_ttc_customers(ws, manager, results["roster"])
                results["customer_records"][manager][sheet] = records
                results["flagged_customers"][manager][sheet] = filter_significant_customers(
                    records, sheet, TTC_ALERT_THRESHOLDS
                )
    except Exception as e:
        results["errors"].append(f"Customer file error: {str(e)}")

    if monthly_path:
        results["monthly"] = analyze_ttc_monthly(monthly_path, results["roster"])

    return results
